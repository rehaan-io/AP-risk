import io
import docx
import openpyxl
from services.azure_document_intelligence import analyze_policy_bytes
from services.llm_service import llm_provider

async def extract_and_interpret_policy(filename: str, file_bytes: bytes):
    """
    Extracts text based on file format, then calls Gemini to interpret it.
    """
    ext = filename.split('.')[-1].lower()
    extracted_text = ""
    
    try:
        if ext == 'txt':
            extracted_text = file_bytes.decode('utf-8', errors='ignore')
            
        elif ext == 'docx':
            doc = docx.Document(io.BytesIO(file_bytes))
            lines = []
            for para in doc.paragraphs:
                if para.text.strip():
                    lines.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        lines.append(" | ".join(row_data))
            extracted_text = "\n".join(lines)
            
        elif ext == 'xlsx':
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            text_lines = []
            for sheet in wb.worksheets:
                text_lines.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(v) for v in row if v is not None])
                    if row_text:
                        text_lines.append(row_text)
            extracted_text = "\n".join(text_lines)
            
        elif ext in ['pdf', 'png', 'jpg', 'jpeg']:
            # Use Azure Document Intelligence for PDFs and images
            extracted_text = await analyze_policy_bytes(file_bytes)
            
        else:
            raise ValueError(f"Unsupported policy format: {ext}")
            
    except Exception as e:
        print(f"Error extracting content from {filename}: {e}")
        return "", []
        
    # Send extracted text to LLM
    try:
        rules = await llm_provider.interpret_policy(extracted_text)
    except Exception as e:
        print(f"Extraction completely failed after retries: {e}")
        rules = [{'rule_type': 'POLICY_EXTRACTION_FAILED', 'status': 'REQUIRES_CONFIRMATION', 'confidence': 0.0, 'reason': str(e)}]
            
    return extracted_text, rules
